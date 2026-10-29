"""Build a sanitized report and presentation addendum from the supplied MevzuAI outputs."""

from html import escape
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE_PRESENTATION = ROOT / "data/temp_reference/MevzuAI_Sunum.pptx"
SOURCE_REPORTS = (
    ROOT / "data/temp_reference/birlesik_mevzuat_uyum_risk_raporu_20261024_1710.html",
    ROOT / "data/temp_reference/remediation_farklari_top5_tur2_20261024_2000.html",
)
OUTPUT_PRESENTATION = ROOT / "docs/MevzuAI_Sunum_Guncel.pptx"
OUTPUT_MARKDOWN = ROOT / "docs/MevzuAI_Birlesik_Risk_ve_Remediation_Ozeti.md"
OUTPUT_HTML = ROOT / "docs/MevzuAI_Birlesik_Risk_ve_Remediation_Ozeti.html"
ASSETS_DIR = ROOT / "docs/assets/mevuai_evidence"

NAVY = RGBColor(17, 34, 64)
NAVY_LIGHT = RGBColor(29, 54, 91)
TEAL = RGBColor(0, 175, 160)
CYAN = RGBColor(53, 190, 215)
GOLD = RGBColor(248, 188, 74)
RED = RGBColor(214, 87, 94)
INK = RGBColor(34, 48, 69)
MUTED = RGBColor(99, 115, 137)
PAPER = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(231, 240, 250)
PALE_GOLD = RGBColor(255, 246, 223)

CAPTURE_WIDTH = 1800
CAPTURE_HEIGHT = 740
CAPTURE_BG = "#101827"
CAPTURE_PANEL = "#182437"
CAPTURE_MUTED = "#9aabc1"
CAPTURE_WHITE = "#eef4fb"
CAPTURE_GREEN = "#71d7a1"
CAPTURE_RED = "#ff8989"
CAPTURE_CYAN = "#71d8e6"
CAPTURE_GOLD = "#ffd37a"


FINDINGS = [
    ("KRİTİK", 24),
    ("YÜKSEK", 23),
    ("ORTA", 21),
    ("DÜŞÜK", 9),
    ("BİLGİ", 5),
]

REMEDIATIONS = [
    (
        "RISK-C05",
        "Safir.Abex · Kimlik bilgisi",
        "Kod içindeki Basic-auth sabitini kaldırıp yapılandırma/secret store "
        "üzerinden okumak; mevcut credential rotasyonu ayrıca izlenmeli.",
    ),
    (
        "RISK-C08",
        "safir-online-sobe · Eşleşme ve log",
        "ZERO_CVV merchant eşleşmesini tam eşleşmeye çevirmek; PIN/F55 benzeri "
        "hassas trace alanlarını maskelemek.",
    ),
    (
        "RISK-C09",
        "acqua-dcc-monex · DCC teklif akışı",
        "İstek kimliğiyle idempotency sağlamak ve geçersiz kur değerlerini "
        "POS'a dönmeden önce reddetmek.",
    ),
    (
        "RISK-C10",
        "acqua-cassandra-data-archiver · Arşiv",
        "S3 yüklemesini nesne metadata/boyut doğrulamasıyla teyit etmeden "
        "kaynak kayıtlarını silmemek.",
    ),
    (
        "RISK-C11",
        "acqua-rrn-generator · Sayaç",
        "Cassandra LWT [applied] sonucunu denetlemek; başarısızlıkta sınırlı "
        "yeniden deneme ve açık hata üretmek.",
    ),
]


def _markdown() -> str:
    distribution = " · ".join(f"{count} {severity}" for severity, count in FINDINGS)
    lines = [
        "# MevzuAI — Birleşik Risk ve Remediation Özeti",
        "",
        "> Kaynak: 24.10.2026 tarihli birleşik mevzuat uyum/risk raporu ve "
        "remediation tur 2 fark raporu. Bu belge, verilen raporların güvenli "
        "yönetici özetidir; bağımsız yeniden tarama değildir.",
        "",
        "## Yönetici özeti",
        "",
        "| Ölçüm | Kaynak rapordaki sonuç |",
        "|---|---:|",
        "| Taranan depo | 35 |",
        "| Toplam bulgu | 82 |",
        f"| Bulgu dağılımı | {distribution} |",
        "| Depo bazında en yüksek risk | 17 kritik · 8 yüksek · 10 orta |",
        "| Mevzuat uyum değerlendirmesi | %67 (61 değerlendirme) |",
        "| Değerlendirme kırılımı | 41 uyumlu · 10 kısmi · 10 ihlal riski |",
        "| Doğrulanmalı etiketi | 9 |",
        "| Kaynak dayanağı | 13 RAG chunk · 9 DB-RAG atfı |",
        "",
        "Kapsam raporda Kanun 5464 ve Yönetmelik 11180 olarak belirtilmiştir. "
        "Önceki raporlar girdi alınmadan taze tarama yapıldığı raporlanmıştır.",
        "",
        "## Remediation tur 2",
        "",
        "Kullanıcı onayıyla beş kritik bulgu seçilmiştir. Fark raporu toplam "
        "**5 depo, 9 dosya, +89/-10 satır** değişikliği bildirir.",
        "",
        "| Bulgu | Alan | Özetlenen değişiklik |",
        "|---|---|---|",
    ]
    for finding_id, area, change in REMEDIATIONS:
        lines.append(f"| {finding_id} | {area} | {change} |")
    lines.extend(
        [
            "",
            "## Doğrulama ve açık işler",
            "",
            "- Fark raporunda 3 konu **Doğrulanmalı** olarak işaretlenmiştir; "
            "iş kararı, ortam konfigürasyonu veya istemci davranışı ek doğrulama ister.",
            "- C06 credential rotasyonu/vault işi; C12/C13 altyapı-konfigürasyon "
            "işi olarak bu turda kapsam dışı bırakılmıştır. C14/C15 sonraki tura "
            "ertelenmiştir.",
            "- Kaynak fark raporu bir PR dalı adı verir; ancak bu özet, testlerin "
            "geçtiğini veya GitHub PR'ın açıldığını doğrulamamaktadır.",
            "- Kaynak rapordaki credential/token örnekleri bu özet ve sunuma "
            "aktarılmamıştır. Credential rotasyonu ve Git geçmişi temizliği "
            "ayrı operasyonel işlerdir.",
            "",
            "## Yorumlama sınırı",
            "",
            "Rapor sayıları ve remediation açıklamaları verilen HTML belgelerinden "
            "aktarılmıştır; bu çalışma kaynak depoları yeniden taramadı, kod "
            "farklarını bağımsız test etmedi ve hukuki uygunluk görüşü sunmaz. "
            "Düzeltmeler merge edilmeden önce ilgili depo sahipleri, güvenlik "
            "ekibi ve mevzuat uzmanı tarafından doğrulanmalıdır.",
            "",
        ]
    )
    return "\n".join(lines)


def _html_report() -> str:
    finding_rows = "".join(
        f"<tr><th>{escape(finding_id)}</th><td>{escape(area)}</td>"
        f"<td>{escape(change)}</td></tr>"
        for finding_id, area, change in REMEDIATIONS
    )
    finding_distribution = "".join(
        f'<div class="distribution"><span>{escape(severity)}</span>'
        f"<strong>{count}</strong></div>"
        for severity, count in FINDINGS
    )
    return f"""<!doctype html>
<html lang="tr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>MevzuAI — Birleşik Risk ve Remediation Özeti</title>
<style>
:root{{--navy:#112240;--teal:#00afa0;--ink:#223045;--muted:#637389;
--paper:#f7f9fc;--line:#dee5ee;--blue:#e7f0fa;--gold:#fff6df;}}
*{{box-sizing:border-box}} body{{margin:0;background:var(--paper);color:var(--ink);
font:15px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif}}
main{{max-width:1120px;margin:36px auto;padding:0 22px 48px}}
header{{background:var(--navy);color:white;padding:34px 38px;border-radius:18px}}
.eyebrow{{color:#35bed7;font-weight:700;letter-spacing:.1em;text-transform:uppercase}}
h1{{font-size:30px;line-height:1.2;margin:8px 0}}
header p{{color:#d7e3f0;margin:8px 0 0}}
h2{{color:var(--navy);font-size:21px;margin:30px 0 12px}}
.cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:14px;margin-top:20px}}
.card{{background:white;border:1px solid var(--line);border-radius:14px;padding:18px}}
.card strong{{display:block;font-size:30px;color:var(--navy)}}
.card span{{color:var(--muted);font-size:13px}}
.distribution-grid{{display:grid;grid-template-columns:repeat(5,1fr);gap:10px}}
.distribution{{background:white;border:1px solid var(--line);padding:12px 14px;
border-radius:10px;display:flex;justify-content:space-between}}
.distribution span{{font-size:13px;color:var(--muted)}}
table{{width:100%;border-collapse:collapse;background:white;border-radius:12px;
overflow:hidden;border:1px solid var(--line)}}
th,td{{padding:13px 14px;text-align:left;vertical-align:top;border-bottom:1px solid var(--line)}}
thead th{{background:var(--navy);color:white}} tbody th{{white-space:nowrap;color:var(--navy)}}
.notice{{background:var(--gold);border-left:4px solid #f8bc4a;padding:15px 18px;
border-radius:8px}} .source{{color:var(--muted);font-size:13px}}
ul{{padding-left:22px}} li{{margin:7px 0}}
@media(max-width:760px){{.cards{{grid-template-columns:repeat(2,1fr)}}
.distribution-grid{{grid-template-columns:repeat(2,1fr)}} header{{padding:25px}}}}
@media print{{body{{background:white}}main{{margin:0;max-width:none}}header{{border-radius:0}}}}
</style>
</head>
<body><main>
<header>
<div class="eyebrow">Yönetici özeti · 24 Ekim 2026</div>
<h1>Birleşik mevzuat uyumu ve depo risk analizi</h1>
<p>Kanun 5464 ve Yönetmelik 11180 kapsamındaki kaynak raporların güvenli özeti</p>
</header>
<section class="cards" aria-label="Özet ölçümler">
<div class="card"><strong>35</strong><span>Taranan depo</span></div>
<div class="card"><strong>82</strong><span>Toplam bulgu</span></div>
<div class="card"><strong>%67</strong><span>61 değerlendirmede uyum oranı</span></div>
<div class="card"><strong>9</strong><span>Doğrulanmalı etiketi</span></div>
</section>
<h2>Bulgu dağılımı</h2>
<div class="distribution-grid">{finding_distribution}</div>
<h2>Mevzuat değerlendirmesi</h2>
<p>61 hüküm-depo değerlendirmesinin 41'i uyumlu, 10'u kısmi ve 10'u ihlal riski
olarak raporlanmıştır. Depo bazında en yüksek risk dağılımı 17 kritik, 8 yüksek
ve 10 orta seviyedir. Kaynakta 13 RAG chunk ve 9 DB-RAG atfı kullanıldığı belirtilir.</p>
<h2>Remediation tur 2 — onaylanan beş kritik bulgu</h2>
<p>Fark raporu 5 depo, 9 dosya ve toplam +89/-10 satır değişikliği bildirir.</p>
<table><thead><tr><th>Bulgu</th><th>Alan</th><th>Düzeltme özeti</th></tr></thead>
<tbody>{finding_rows}</tbody></table>
<h2>Açık doğrulama ve operasyon işleri</h2>
<ul>
<li>Fark raporunda 3 konu “Doğrulanmalı” olarak işaretlenmiştir.</li>
<li>Credential rotasyonu/vault ve altyapı-konfigürasyon işleri kod farkından
ayrı takip edilmelidir; sonraki tur bulguları bu kapsamda değildir.</li>
<li>Kaynak rapor bir PR dalı belirtir; testlerin geçtiği veya GitHub PR'ın
açıldığı bu belgelerle doğrulanmamıştır.</li>
<li>Credential/token değerleri bu özete aktarılmamıştır. Rotasyon ve Git geçmişi
temizliği ayrı operasyonel takip gerektirir.</li>
</ul>
<p class="notice"><strong>Yorumlama sınırı:</strong> Bu özet, sağlanan HTML
raporlarındaki bilgileri düzenler. Bağımsız kod taraması, test çalıştırması veya
hukuki uygunluk değerlendirmesi değildir. Merge öncesinde depo sahipleri,
güvenlik ekibi ve mevzuat uzmanı doğrulamalıdır.</p>
<p class="source">Kaynaklar: birlesik_mevzuat_uyum_risk_raporu_20261024_1710.html;
remediation_farklari_top5_tur2_20261024_2000.html.</p>
</main></body></html>"""


def _capture_font(size: int, bold: bool = False):
    font_path = (
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
        if bold
        else "/System/Library/Fonts/Supplemental/Arial.ttf"
    )
    try:
        return ImageFont.truetype(font_path, size)
    except OSError:
        return ImageFont.load_default(size=size)


def _create_capture(
    filename: str,
    title: str,
    subtitle: str,
    lines: list[tuple[str, str]],
    footer: str,
) -> Path:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    output = ASSETS_DIR / filename
    image = Image.new("RGB", (CAPTURE_WIDTH, CAPTURE_HEIGHT), CAPTURE_BG)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle(
        (24, 20, CAPTURE_WIDTH - 24, CAPTURE_HEIGHT - 20),
        radius=22,
        fill=CAPTURE_PANEL,
        outline="#35445a",
        width=2,
    )
    for x, color in ((54, "#ff6b6b"), (88, "#ffc857"), (122, "#5bd49b")):
        draw.ellipse((x, 43, x + 17, 60), fill=color)

    draw.text((168, 37), "MEVZUATAGENT · KANIT GÖRSELİ", font=_capture_font(22, True), fill=CAPTURE_MUTED)
    draw.text((60, 100), title, font=_capture_font(36, True), fill=CAPTURE_WHITE)
    draw.text((62, 153), subtitle, font=_capture_font(22), fill=CAPTURE_CYAN)

    y = 214
    for text, color in lines:
        draw.text((68, y), text, font=_capture_font(22), fill=color)
        y += 38
    draw.line((58, CAPTURE_HEIGHT - 86, CAPTURE_WIDTH - 58, CAPTURE_HEIGHT - 86), fill="#35445a", width=2)
    draw.text(
        (62, CAPTURE_HEIGHT - 70),
        footer,
        font=_capture_font(17),
        fill=CAPTURE_GOLD,
    )
    image.save(output, optimize=True)
    return output


def _create_evidence_captures() -> dict[str, Path]:
    captures = {
        "setup": _create_capture(
            "kurulum_asamalari.png",
            "macOS kurulum akışı",
            "Görsel, install.sh içindeki gerçek adım sırasından oluşturulmuştur",
            [
                ("$ ./install.sh", CAPTURE_GREEN),
                ("01  macOS, Git ve Python 3.10+ ön koşullarını denetle", CAPTURE_WHITE),
                ("02  .venv sanal ortamını oluştur; bağımlılıkları yükle", CAPTURE_WHITE),
                ("03  JSON yapılandırmasını ve çalışma klasörlerini doğrula", CAPTURE_WHITE),
                ("04  Ollama / model kurulumu için ayrı kullanıcı onayı al", CAPTURE_WHITE),
                ("05  Mevzuatı indirip yerel BM25 indeksi oluşturmayı sor", CAPTURE_WHITE),
                ("06  repo_link.json depolarını klonlamak için onay iste", CAPTURE_WHITE),
                ("07  Uygulama ve pytest komutlarını göster", CAPTURE_WHITE),
            ],
            "KAYNAK: install.sh · Şematik ekran görselidir; gerçek kurulum oturumu kaydı değildir.",
        ),
        "repository": _create_capture(
            "sentetik_test_deposu.png",
            "Sentetik test deposu",
            "Git'te izlenen gerçek fixture yapısı · main @ a14df30",
            [
                ("mevzuatagent-test-fixtures/", CAPTURE_CYAN),
                ("├── README.md", CAPTURE_WHITE),
                ("├── demo1/  payment_rules.py · scan_result.json · remediation_plan.md", CAPTURE_WHITE),
                ("│   └── scan_result.md / scan_result.html", CAPTURE_WHITE),
                ("└── demo2/  payment_rules.py · scan_result.json · remediation_plan.md", CAPTURE_WHITE),
                ("    └── scan_result.md / scan_result.html", CAPTURE_WHITE),
                ("2 demo · toplam 10 sentetik bulgu · TEST_FIXTURE_NOT_LAW", CAPTURE_GOLD),
                ("YALNIZCA TEST İÇİNDİR — gerçek mevzuat tespiti değildir.", CAPTURE_RED),
            ],
            "KAYNAK: data/test_repositories/mevzuatagent-test-fixtures · Git ls-files / README",
        ),
        "diff": _create_capture(
            "remediation_diff_ve_test.png",
            "Remediation diff ve kabul testi",
            "Gerçek fixture raporundan alınan örnek değişiklikler · test: 2 passed",
            [
                ("demo1/payment_rules.py", CAPTURE_CYAN),
                ("-    return value", CAPTURE_RED),
                ("+    return '*' * max(0, len(value) - visible) + value[-visible:]", CAPTURE_GREEN),
                ("-    return amount >= 0", CAPTURE_RED),
                ("+    return amount > 0", CAPTURE_GREEN),
                ("$ python -m pytest -q", CAPTURE_MUTED),
                ("..                                                                       [100%]", CAPTURE_GREEN),
                ("2 passed · fixture kabul testleri · yalnız test amaçlı", CAPTURE_GOLD),
            ],
            "KAYNAK: data/e2e/reports/remediation + data/e2e/reports/tests",
        ),
        "pr": _create_capture(
            "pr_akisi_yerel_prova.png",
            "PR akışı — yerel prova",
            "GitHub ekranı değildir · gerçek PR oluşturulmamıştır",
            [
                ("$ git --git-dir=data/e2e/pr-origin.git branch -a", CAPTURE_MUTED),
                ("  main", CAPTURE_WHITE),
                ("* mevzuatagent/remediation-c4b060c129", CAPTURE_CYAN),
                ("$ gh pr create  # fake-bin/gh test ikilisi", CAPTURE_MUTED),
                ("PR CLI çağrısı yerel test stub'ına yönlendirildi", CAPTURE_WHITE),
                ("Yerel bare remote / sahte gh ile entegrasyon provası", CAPTURE_GREEN),
                ("Gerçek GitHub PR: OLUŞTURULMADI", CAPTURE_RED),
                ("Fixture testidir; GitHub'da PR numarası/URL'si yoktur.", CAPTURE_GOLD),
            ],
            "KAYNAK: data/e2e/pr-origin.git + data/e2e/fake-bin/gh · canlı GitHub işlemi yok",
        ),
    }
    return captures


def _set_shape_text(shape, text: str, size: int, color: RGBColor, bold: bool = False):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color


def _add_text(slide, text, x, y, width, height, size, color=INK, bold=False):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    _set_shape_text(shape, text, size, color, bold)
    return shape


def _add_box(slide, x, y, width, height, fill):
    shape = slide.shapes.add_shape(
        1, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def _new_slide(presentation, eyebrow, title, subtitle):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = PAPER
    _add_text(slide, eyebrow, 0.65, 0.3, 12.0, 0.25, 10, TEAL, True)
    _add_text(slide, title, 0.65, 0.65, 12.0, 0.55, 24, NAVY, True)
    _add_text(slide, subtitle, 0.68, 1.22, 12.0, 0.42, 11, MUTED)
    return slide


def _add_report_slides(presentation):
    slide = _new_slide(
        presentation,
        "KAYNAK RAPORU · 24.10.2026",
        "Birleşik risk görünümü: 35 depo, 82 bulgu",
        "Kanun 5464 + Yönetmelik 11180 · Sağlanan birleşik HTML raporundaki ölçümler",
    )
    metrics = [
        ("35", "taranan depo", TEAL),
        ("82", "toplam bulgu", RED),
        ("%67", "uyum · 61 değerlendirme", CYAN),
        ("9", "Doğrulanmalı", GOLD),
    ]
    for index, (value, label, accent) in enumerate(metrics):
        x = 0.68 + index * 3.12
        _add_box(slide, x, 1.9, 2.86, 1.15, WHITE)
        _add_box(slide, x, 1.9, 0.08, 1.15, accent)
        _add_text(slide, value, x + 0.2, 2.02, 2.45, 0.48, 25, NAVY, True)
        _add_text(slide, label, x + 0.2, 2.55, 2.45, 0.3, 10, MUTED)

    distribution = "   ·   ".join(
        f"{count} {severity}" for severity, count in FINDINGS
    )
    _add_text(slide, "Bulgu dağılımı", 0.78, 3.4, 2.0, 0.3, 14, NAVY, True)
    _add_text(slide, distribution, 0.78, 3.78, 11.8, 0.42, 13, INK, True)
    _add_box(slide, 0.68, 4.55, 12.0, 1.35, PALE_BLUE)
    _add_text(
        slide,
        "61 hüküm-depo değerlendirmesi: 41 uyumlu · 10 kısmi · 10 ihlal riski\n"
        "Depo bazında en yüksek risk: 17 kritik · 8 yüksek · 10 orta\n"
        "Kaynak dayanağı: 13 RAG chunk + 9 DB-RAG atfı",
        0.93,
        4.73,
        11.45,
        0.95,
        13,
        NAVY,
    )
    _add_text(
        slide,
        "Kaynak rapor ölçümleri; bağımsız yeniden tarama veya hukuki görüş değildir.",
        0.75,
        6.35,
        11.9,
        0.3,
        10,
        MUTED,
    )

    slide = _new_slide(
        presentation,
        "REMEDIATION · TUR 2",
        "Kullanıcı onayıyla seçilen beş kritik bulgu",
        "Fark raporu: 5 depo · 9 dosya · +89 / -10 satır · test/PR sonucu bu rapordan doğrulanmıyor",
    )
    table_shape = slide.shapes.add_table(
        6, 3, Inches(0.68), Inches(1.85), Inches(12.0), Inches(4.55)
    )
    table = table_shape.table
    table.columns[0].width = Inches(1.2)
    table.columns[1].width = Inches(3.2)
    table.columns[2].width = Inches(7.6)
    rows = [("Bulgu", "Alan", "Fark raporundaki düzeltme özeti"), *REMEDIATIONS]
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if row_index == 0 else (
                WHITE if row_index % 2 else PALE_BLUE
            )
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(11 if row_index else 12)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = WHITE if row_index == 0 else INK
    _add_text(
        slide,
        "Credential/token örnekleri sunuma aktarılmadı; rotasyon ve Git geçmişi temizliği ayrı operasyonel işlerdir.",
        0.75,
        6.58,
        12.0,
        0.3,
        9,
        MUTED,
    )

    slide = _new_slide(
        presentation,
        "KAPANIŞ · DOĞRULAMA DURUMU",
        "Fark raporu sonuç kanıtı değildir",
        "Kod farkı, test sonucu, canlı PR ve operasyonel riskler ayrı ayrı doğrulanmalıdır.",
    )
    cards = [
        (
            "Açık doğrulama",
            "3 madde kaynak raporda “Doğrulanmalı”.\n"
            "İş kararı, ortam ayarı veya istemci davranışı netleştirilmeli.",
            CYAN,
        ),
        (
            "Ayrı operasyon işleri",
            "Credential rotasyonu / vault ve altyapı-konfigürasyon işleri.\n"
            "Kod farkı bu işleri tek başına kapatmaz.",
            GOLD,
        ),
        (
            "Sonraki kapılar",
            "Testleri çalıştır ve sonuçları kaydet.\n"
            "Diff'i gözden geçir; gerçek PR durumunu GitHub'da doğrula.",
            TEAL,
        ),
    ]
    for index, (title, body, accent) in enumerate(cards):
        x = 0.68 + index * 4.05
        _add_box(slide, x, 2.0, 3.78, 2.3, WHITE)
        _add_box(slide, x, 2.0, 3.78, 0.1, accent)
        _add_text(slide, title, x + 0.18, 2.25, 3.4, 0.42, 15, NAVY, True)
        _add_text(slide, body, x + 0.18, 2.82, 3.4, 1.2, 11, INK)
    _add_box(slide, 0.68, 4.78, 12.0, 1.0, PALE_GOLD)
    _add_text(
        slide,
        "Sınır: Bu ek, verilen raporları özetler. Kaynak depolar yeniden taranmadı; "
        "testlerin geçtiği veya PR'ın açıldığı iddia edilmez.",
        0.95,
        4.98,
        11.45,
        0.6,
        12,
        NAVY,
        True,
    )
    _add_text(
        slide,
        "Güvenli paylaşım için credential/token değerleri eklerden çıkarılmıştır.",
        0.75,
        6.35,
        11.9,
        0.3,
        10,
        MUTED,
    )


def _add_evidence_slides(presentation, captures: dict[str, Path]):
    slides = [
        (
            "KURULUM · AŞAMA GÖRSELİ",
            "Kurulum aşamaları ve açık onay kapıları",
            "install.sh kaynak kodundan sıralandı; bu makinede yapılmış bir kurulum kaydı olarak sunulmaz.",
            "setup",
            "Görsel belge akışını anlatır; Ollama, model, mevzuat indirme ve klonlama adımları ayrıca onay ister.",
        ),
        (
            "TEST REPOSU · GERÇEK DOSYA YAPISI",
            "İki sentetik demo, on test bulgusu",
            "Depo ağacı mevcut test fixture Git deposundan; içerik yalnızca test için hazırlanmıştır.",
            "repository",
            "Sentetik fixture; gerçek sistem veya gerçek mevzuat bulgusu değildir.",
        ),
        (
            "REMEDIATION · DIFF + TEST",
            "Önce / sonra farkı ve kabul testi",
            "Görsel, izole remediation worktree raporundaki iki örnek değişikliği ve test raporunu gösterir.",
            "diff",
            "Test raporunda 2 sentetik kabul testi geçti; hukuki uyumluluk doğrulaması değildir.",
        ),
        (
            "PR · YEREL ENTEGRASYON PROVASI",
            "Push/PR hattı güvenli test ortamında denendi",
            "Yerel bare Git remote ve fake gh kullanıldı; gerçek GitHub PR açılmadı.",
            "pr",
            "Gerçek PR ekran görüntüsü mevcut değildir; bu görsel yerel mock provayı açıkça etiketler.",
        ),
    ]
    for eyebrow, title, subtitle, key, note in slides:
        slide = _new_slide(presentation, eyebrow, title, subtitle)
        slide.shapes.add_picture(
            str(captures[key]),
            Inches(0.68),
            Inches(1.82),
            width=Inches(12.0),
            height=Inches(4.93),
        )
        _add_text(
            slide,
            note,
            0.75,
            6.85,
            12.0,
            0.28,
            9,
            MUTED,
        )


def main():
    missing_sources = [
        path for path in (SOURCE_PRESENTATION, *SOURCE_REPORTS) if not path.is_file()
    ]
    if missing_sources:
        raise FileNotFoundError(
            "Kaynak belge(ler) bulunamadı; üç dosyayı data/temp_reference/ içine "
            "kopyalayın: " + ", ".join(path.name for path in missing_sources)
        )
    presentation = Presentation(str(SOURCE_PRESENTATION))
    _add_report_slides(presentation)
    captures = _create_evidence_captures()
    _add_evidence_slides(presentation, captures)
    OUTPUT_PRESENTATION.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(OUTPUT_PRESENTATION))

    OUTPUT_MARKDOWN.write_text(_markdown(), encoding="utf-8")
    OUTPUT_HTML.write_text(_html_report(), encoding="utf-8")
    print(f"Sunum: {OUTPUT_PRESENTATION}")
    print(f"Markdown rapor: {OUTPUT_MARKDOWN}")
    print(f"HTML rapor: {OUTPUT_HTML}")


if __name__ == "__main__":
    main()
